#!/usr/bin/env python3
"""실제 모델·OpenProxy에서 지속 업로드/검색 및 큰 DOCX 파일을 측정한다.

DATABASE_URL=... EMBEDDING_PROVIDER=local backend/.venv/bin/python \
  scripts/benchmark_pipeline.py --out benchmark-results
4개 임시 계정과 이 계정의 문서만 만들고, 전용 API/워커를 종료한 뒤 모두 삭제한다.
지정한 API 포트는 비어 있어야 한다. JSON에 접속 정보·계정 비밀번호는 기록하지 않는다.
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import math
import os
import secrets
import socket
import statistics
import sys
import time
import uuid
from contextlib import ExitStack
from io import BytesIO
from pathlib import Path

import httpx
import psycopg
from docx import Document
from docx.shared import Inches
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from openarchive.config import get_settings
from openarchive.services.auth import create_user


def request_summary(rows):
    values = sorted(r["ms"] for r in rows)
    return {
        "requests": len(rows),
        "failures": sum(r.get("status", 500) >= 400 for r in rows),
        "empty_results": sum(bool(r.get("empty_results")) for r in rows),
        "median_ms": statistics.median(values) if values else None,
        "p95_ms": values[int(0.95 * (len(values) - 1))] if values else None,
    }


def upload_schedule(duration, rate):
    return [i / rate for i in range(duration * rate)]


def large_file(source, size):
    data = bytearray()
    number = 0
    while True:
        number += 1
        section = f"\n\nbenchmark section {number}\n\n{source}".encode()
        if len(data) + len(section) > size:
            data.extend(b" " * (size - len(data)))
            return bytes(data)
        data.extend(section)


def large_docx(source, size):
    """원본 크기와 추출 텍스트 길이를 분리한 유효 DOCX를 만든다."""
    document = Document()
    for paragraph in source[:250_000].split("\n\n"):
        document.add_paragraph(paragraph)
    side = math.isqrt(max(1, size // 3))
    image = Image.frombytes("RGB", (side, side), os.urandom(side * side * 3))
    picture = BytesIO()
    image.save(picture, format="PNG", compress_level=0)
    picture.seek(0)
    document.add_picture(picture, width=Inches(6))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


async def run(args):
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", args.port))
    dsn = get_settings().database_url
    run_id = "pipeline-bench-" + uuid.uuid4().hex[:12]
    owners = [f"{run_id}-{i}" for i in range(4)]
    passwords = [secrets.token_urlsafe(24) for _ in owners]
    source = "\n\n".join(
        (ROOT / f).read_text() for f in ["README.md", "docs/ARCHITECTURE.md"]
    )
    report = {
        "model": "BAAI/bge-m3",
        "run_id": run_id,
        "source_sha256": hashlib.sha256(source.encode()).hexdigest(),
        "duration": args.duration,
        "upload_rate": args.upload_rate,
        "accounts": 4,
        "workers": args.workers,
        "requests": [],
        "samples": [],
        "scope": "Mac native API/workers, database via DATABASE_URL",
    }
    procs, clients = [], []
    log_files = ExitStack()
    start = time.perf_counter()
    stage = "warmup"
    stop_monitor = asyncio.Event()
    monitor = None
    args.out.mkdir(parents=True, exist_ok=True)

    def save():
        (args.out / "results.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n"
        )

    async def snapshot():
        async with await psycopg.AsyncConnection.connect(dsn) as conn:
            docs = await (
                await conn.execute(
                    "SELECT count(*),count(*) FILTER (WHERE embedding_status<>'ready') "
                    "FROM documents WHERE owner_id=ANY(%s)",
                    (owners,),
                )
            ).fetchone()
            jobs = await (
                await conn.execute(
                    "SELECT j.kind,j.status,count(*),max(j.attempts) FROM embedding_jobs j "
                    "JOIN documents d ON d.id=j.document_id WHERE d.owner_id=ANY(%s) GROUP BY j.kind,j.status",
                    (owners,),
                )
            ).fetchall()
            chunks, mismatch = await (
                await conn.execute(
                    "SELECT count(*),count(*) FILTER (WHERE c.version<>d.version) FROM document_chunks c "
                    "JOIN documents d ON d.id=c.document_id WHERE d.owner_id=ANY(%s)",
                    (owners,),
                )
            ).fetchone()
        return {
            "t": time.perf_counter() - start,
            "documents": docs[0],
            "not_ready": docs[1],
            "chunks": chunks,
            "mismatch": mismatch,
            "jobs": jobs,
            "active": sum(r[2] for r in jobs if r[1] in ("pending", "processing")),
            "errors": sum(r[2] for r in jobs if r[1] == "error"),
        }

    async def watch():
        while not stop_monitor.is_set():
            report["samples"].append(await snapshot())
            save()
            await asyncio.sleep(1)

    async def settle(seconds):
        since = time.perf_counter()
        while time.perf_counter() - since < seconds:
            status = await snapshot()
            if status["active"] == 0:
                return {
                    "seconds": time.perf_counter() - since,
                    "status": status,
                    "converged": status["errors"]
                    == status["not_ready"]
                    == status["mismatch"]
                    == 0,
                }
            await asyncio.sleep(2)
        return {
            "seconds": time.perf_counter() - since,
            "status": await snapshot(),
            "converged": False,
        }

    async def send(
        index, kind, seq=0, planned=None, file=None, filename="project-large.docx"
    ):
        began = time.perf_counter()
        try:
            headers = {"Idempotency-Key": str(uuid.uuid4())}
            if file is not None:
                response = await clients[index].post(
                    "/documents",
                    files={"file": (filename, file, "application/octet-stream")},
                    data={
                        "title": owners[index] + "-large",
                        "tags": run_id,
                        "visibility": "private",
                    },
                    headers=headers,
                )
            elif kind == "upload":
                offset = (seq * 600) % max(1, len(source) - 600)
                response = await clients[index].post(
                    "/documents/text",
                    headers=headers,
                    json={
                        "title": f"{owners[index]}-{seq}",
                        "content": source[offset : offset + 600] + f"\n측정 {seq}",
                        "tags": [run_id],
                        "visibility": "private",
                    },
                )
            else:
                response = await clients[index].post(
                    "/search",
                    json={"query": "문서 버전과 벡터 정합성", "tags": [run_id], "k": 5},
                )
                if response.status_code == 200:
                    assert all(
                        item["title"].startswith(owners[index])
                        for item in response.json()["items"]
                    ), "private document leak"
            row = {
                "kind": kind,
                "phase": stage,
                "account": index,
                "status": response.status_code,
                "empty_results": kind == "search"
                and response.status_code == 200
                and not response.json()["items"],
                "ms": (time.perf_counter() - began) * 1000,
            }
            if response.status_code >= 400:
                row["detail"] = response.text[:500]
            if planned is not None:
                row["schedule_delay_ms"] = max(0, began - planned) * 1000
            report["requests"].append(row)
            return response
        except httpx.HTTPError as error:
            report["requests"].append(
                {
                    "kind": kind,
                    "phase": stage,
                    "account": index,
                    "error": type(error).__name__,
                    "ms": (time.perf_counter() - began) * 1000,
                }
            )
            return None

    try:
        async with await psycopg.AsyncConnection.connect(dsn) as conn:
            for owner, password in zip(owners, passwords, strict=True):
                await create_user(conn, owner, password)
        env = os.environ | {
            "DATABASE_URL": dsn,
            "EMBEDDING_PROVIDER": "local",
            "TOKENIZERS_PARALLELISM": "false",
        }
        for name, command in [
            (
                "api",
                [
                    "-m",
                    "uvicorn",
                    "openarchive.main:app",
                    "--host",
                    "127.0.0.1",
                    "--port",
                    str(args.port),
                ],
            ),
            *[
                (
                    "worker" if i == 0 else f"worker-{i + 1}",
                    ["-m", "openarchive.worker"],
                )
                for i in range(args.workers)
            ],
        ]:
            handle = log_files.enter_context(
                await asyncio.to_thread(Path.open, args.out / f"{name}.log", "w")
            )
            procs.append(
                await asyncio.create_subprocess_exec(
                    sys.executable,
                    *command,
                    cwd=ROOT / "backend",
                    env=env,
                    stdout=handle,
                    stderr=asyncio.subprocess.STDOUT,
                )
            )
        clients = [
            httpx.AsyncClient(
                base_url=f"http://127.0.0.1:{args.port}/api/", timeout=120
            )
            for _ in owners
        ]
        for _ in range(45):
            try:
                if (await clients[0].get("/health")).status_code == 200:
                    break
            except httpx.HTTPError:
                pass
            await asyncio.sleep(1)
        for index, (owner, password) in enumerate(zip(owners, passwords, strict=True)):
            (
                await clients[index].post(
                    "/auth/login", json={"username": owner, "password": password}
                )
            ).raise_for_status()
            (await send(index, "upload", seq=-1)).raise_for_status()
        report["warmup"] = await settle(180)
        if not report["warmup"]["converged"]:
            raise RuntimeError("검증 문서의 검색 준비가 끝나지 않았습니다.")
        monitor = asyncio.create_task(watch())
        stage = "steady"
        origin = time.perf_counter()
        requests = []

        async def upload_at(seq, offset):
            await asyncio.sleep(max(0, origin + offset - time.perf_counter()))
            return await send(seq % 4, "upload", seq, planned=origin + offset)

        async def searches(index):
            while time.perf_counter() - origin < args.duration:
                await send(index, "search")
                await asyncio.sleep(0.25)

        requests = [
            asyncio.create_task(upload_at(i, t))
            for i, t in enumerate(upload_schedule(args.duration, args.upload_rate))
        ]
        await asyncio.gather(*requests, *(searches(i) for i in range(4)))
        report["arrival_seconds"] = time.perf_counter() - origin
        report["after_arrivals"] = await snapshot()
        report["steady_settle"] = await settle(300)
        print(
            "Sustained arrivals finished:",
            report["after_arrivals"],
            report["steady_settle"],
            flush=True,
        )
        stage = "large_file"
        guard = await send(
            0,
            "text_limit_guard",
            file=large_file("OpenArchive benchmark text", 500_001),
            filename="overlong-project.txt",
        )
        report["text_limit_guard"] = {
            "status": guard.status_code if guard is not None else None,
            "expected": 400,
        }
        async with await psycopg.AsyncConnection.connect(dsn) as conn:
            report["text_limit_guard"]["saved_documents"] = (
                await (
                    await conn.execute(
                        "SELECT count(*) FROM documents WHERE owner_id=%s AND title=%s",
                        (owners[0], owners[0] + "-large"),
                    )
                ).fetchone()
            )[0]
        data = await asyncio.to_thread(large_docx, source, args.file_mb * 1_000_000)
        response = await send(0, "large_file", file=data)
        report["large_file"] = {
            "bytes": len(data),
            "status": response.status_code if response is not None else None,
        }
        # A small document arrives behind the large job; this reveals head-of-line waiting.
        small = await send(1, "upload", seq=100000)
        if small is not None and small.status_code == 201:
            report["following_small_document_id"] = small.json()["id"]
        probe_stop = asyncio.Event()

        async def search_probe():
            while not probe_stop.is_set():
                await send(1, "search")
                await asyncio.sleep(1)

        probe_task = asyncio.create_task(search_probe())
        try:
            report["large_settle"] = await settle(args.settle)
        finally:
            probe_stop.set()
            await probe_task
        async with await psycopg.AsyncConnection.connect(dsn) as conn:
            report["job_results"] = await (
                await conn.execute(
                    "SELECT j.kind,j.status,j.attempts,j.last_error,extract(epoch FROM(j.finished_at-j.started_at)),extract(epoch FROM(j.started_at-j.created_at)),j.document_id::text "
                    "FROM embedding_jobs j JOIN documents d ON d.id=j.document_id WHERE d.owner_id=ANY(%s) ORDER BY j.id",
                    (owners,),
                )
            ).fetchall()
            if response is not None and response.status_code == 201:
                did = response.json()["id"]
                saved = await (
                    await conn.execute(
                        "SELECT data FROM document_files WHERE document_id=%s", (did,)
                    )
                ).fetchone()
                report["large_file"]["original_hash_matches"] = (
                    saved is not None
                    and hashlib.sha256(bytes(saved[0])).digest()
                    == hashlib.sha256(data).digest()
                )
                report["large_file"]["chunks"] = (
                    await (
                        await conn.execute(
                            "SELECT count(*) FROM document_chunks WHERE document_id=%s",
                            (did,),
                        )
                    ).fetchone()
                )[0]
            if small is not None and small.status_code == 201:
                report["following_small_ready"] = (
                    await (
                        await conn.execute(
                            "SELECT embedding_status FROM documents WHERE id=%s",
                            (small.json()["id"],),
                        )
                    ).fetchone()
                )[0]
        report["summary"] = {
            kind: request_summary([r for r in report["requests"] if r["kind"] == kind])
            for kind in ("upload", "search", "large_file")
        }
        report["phase_summary"] = {
            phase: {
                kind: request_summary(
                    [
                        r
                        for r in report["requests"]
                        if r["kind"] == kind and r["phase"] == phase
                    ]
                )
                for kind in ("upload", "search", "large_file")
            }
            for phase in ("steady", "large_file")
        }
        print(report["summary"], report["large_settle"], flush=True)
    finally:
        stop_monitor.set()
        if monitor is not None:
            await monitor
        for client in clients:
            await client.aclose()
        for process in procs:
            if process.returncode is None:
                process.terminate()
        for process in procs:
            try:
                await asyncio.wait_for(process.wait(), timeout=30)
            except TimeoutError:
                process.kill()
                await process.wait()
        log_files.close()
        async with await psycopg.AsyncConnection.connect(dsn) as conn:
            await conn.execute(
                "DELETE FROM documents WHERE owner_id=ANY(%s)", (owners,)
            )
            await conn.execute("DELETE FROM users WHERE username=ANY(%s)", (owners,))
        report["cleaned_up"] = True
        save()
    return (
        0
        if report["text_limit_guard"]["status"] == 400
        and report["text_limit_guard"]["saved_documents"] == 0
        and report["large_file"].get("original_hash_matches") is True
        and report["large_file"].get("chunks", 0) > 0
        and report.get("following_small_ready") == "ready"
        and report["steady_settle"]["converged"]
        and report["large_settle"]["converged"]
        and all(
            r["failures"] == r["empty_results"] == 0 for r in report["summary"].values()
        )
        else 1
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--duration", type=int, default=30)
    parser.add_argument("--upload-rate", type=int, default=4)
    parser.add_argument("--file-mb", type=int, default=20)
    parser.add_argument("--workers", type=int, default=1)
    parser.add_argument("--settle", type=int, default=900)
    parser.add_argument("--port", type=int, default=18024)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    if (
        min(args.duration, args.upload_rate, args.file_mb, args.settle, args.workers)
        < 1
    ):
        parser.error("개수는 양의 정수여야 합니다")
    return asyncio.run(run(args))


if __name__ == "__main__":
    raise SystemExit(main())
