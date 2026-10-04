#!/usr/bin/env python3
"""형식별 평가셋(#167)의 실문서를 업로드와 같은 진입점으로 적재한다.

    DATABASE_URL=… python scripts/eval_formats_load.py scripts/eval/formats.json [--clean]

정책브리핑 첨부는 공공누리가 텍스트에 한해서라(내장 로고·삽화는 별도 허락) 저장소에 넣지 않는다.
`fetch`가 있는 자료는 없으면 정책브리핑에서 받고, `derive`가 있는 자료는 받은 원본 PDF의 한 쪽을
스캔처럼 가공해 만든다(`scan_jpg`: JPEG, `image_pptx`: 그 그림 한 장뿐인 슬라이드). 받은 파일의
해시가 평가셋과 다르면 멈춘다. 파생 파일은 라이브러리 판에 따라 바이트가 달라질 수 있어 경고만 한다.

모든 파일의 SHA-256을 평가셋과 대조한 뒤, 각 파일을 `create_document`로 `owner`(평가셋
`users`의 계정)의 비공개 문서로 만들고 평가셋 `tag`를 붙인다. `replaced_by`가 있으면 원본
교체(`replace_original_file`)로 새 판을 쌓는다 — 현재/과거 질문은 이 교체가 만든 텍스트 v2를
잰다. 교체가 다음 텍스트 버전을 만들지 않거나, 보관된 원본(`get_original_file`)의 해시가 넣은 파일과
다르면 멈춘다. 추출 텍스트가 비어 업로드가 거부되는 파일은 실패로 멈추지 않고 `rejected`로 기록한다.

OCR 문서는 「추출 중」으로 생긴다. 임베딩·OCR은 워커가 하므로 적재 뒤 워커를 돌려 잡을 비우고
`eval_search.py --user <evaluator> --tag <tag>`로 잰다. 같은 태그 문서가 이미 있으면 거부한다 —
`--clean`이 그 태그의 평가 계정 문서를 지운다(측정 후 정리에도 쓴다).
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import io
import json
import sys
import urllib.request
import zipfile
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
import psycopg
import pypdfium2
from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from openarchive.config import get_settings
from openarchive.services.documents import (
    EmptyExtractedText,
    create_document,
    get_original_file,
    replace_original_file,
)

FIXED_TIME = datetime(2026, 10, 4, tzinfo=UTC)
DOWNLOAD_URL = "https://www.korea.kr/common/download.do?fileId={file_id}&tblKey=GMN"


def download_attachment(file_id: str) -> bytes:
    request = urllib.request.Request(
        DOWNLOAD_URL.format(file_id=file_id), headers={"User-Agent": "Mozilla/5.0"}
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        return response.read()


def derive_scan_jpg(pdf: bytes, *, page: int) -> bytes:
    """PDF 한 쪽을 스캔본처럼 만든다 — 300dpi 래스터 → 150dpi 축소 · 1.5° 기울임 · 노이즈 · 블러 · JPEG q60.

    `backend/tests/fixtures/SOURCE.md`의 스캔 픽스처와 같은 가공이다. 노이즈 시드를 고정해 같은 원본·같은
    라이브러리에서는 같은 바이트가 나온다.
    """
    document = pypdfium2.PdfDocument(pdf)
    try:
        image = document[page - 1].render(scale=300 / 72).to_pil().convert("L")
    finally:
        document.close()
    image = image.resize((image.width // 2, image.height // 2), Image.LANCZOS)
    image = image.rotate(1.5, resample=Image.BICUBIC, expand=True, fillcolor=255)
    noise = np.random.default_rng(167).normal(0, 12, (image.height, image.width))
    pixels = np.clip(np.asarray(image, dtype=np.float32) + noise, 0, 255).astype(np.uint8)
    image = Image.fromarray(pixels).filter(ImageFilter.GaussianBlur(0.6))
    output = io.BytesIO()
    image.save(output, format="JPEG", quality=60, dpi=(150, 150))
    return output.getvalue()


def derive_image_pptx(pdf: bytes, *, page: int) -> bytes:
    """스캔 가공한 쪽 그림 한 장만 담은 슬라이드. 글상자·노트가 없어 추출 텍스트가 비어야 한다."""
    scan = derive_scan_jpg(pdf, page=page)
    width, height = Image.open(io.BytesIO(scan)).size
    presentation = Presentation()
    presentation.slide_width = Emu(width * 9525)
    presentation.slide_height = Emu(height * 9525)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(scan), 0, 0, presentation.slide_width, presentation.slide_height)
    presentation.core_properties.created = presentation.core_properties.modified = FIXED_TIME
    saved = io.BytesIO()
    presentation.save(saved)
    # 저장 시각이 문서 속성과 zip 항목 시각에 들어간다 — 둘 다 고정해야 같은 원본에서 같은 바이트가 나온다
    output = io.BytesIO()
    with zipfile.ZipFile(saved) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0)),
                            source.read(info.filename), compress_type=zipfile.ZIP_DEFLATED)
    return output.getvalue()


_DERIVE = {"scan_jpg": derive_scan_jpg, "image_pptx": derive_image_pptx}


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def source_files(evalset: dict) -> list[dict]:
    """평가셋이 적은 파일 전부 — 각 자료와, 원본 교체가 있으면 그 교체 파일."""
    return [item for source in evalset["sources"] for item in (source, source.get("replaced_by")) if item is not None]


def _fetch(item: dict, root: Path, download: Callable[[str], bytes]) -> bytes:
    path = root / item["path"]
    if not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(download(item["fetch"]["file_id"]))
    data = path.read_bytes()
    if _sha256(data) != item["sha256"]:
        raise ValueError(f"받은 자료의 해시가 평가셋과 다르다: {item['path']}")
    return data


def ensure_sources(
    evalset: dict, root: Path, *, download: Callable[[str], bytes] = download_attachment
) -> list[str]:
    """없는 자료를 받고 파생 자료를 만든다. 파생 파일 해시가 평가셋과 다르면 경고 문구를 돌려준다."""
    warnings = []
    for item in source_files(evalset):
        if "fetch" in item:
            _fetch(item, root, download)
        elif "derive" in item:
            spec = item["derive"]
            data = _DERIVE[spec["kind"]](_fetch(spec["from"], root, download), page=spec["page"])
            path = root / item["path"]
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
            if _sha256(data) != item["sha256"]:
                warnings.append(f"파생 자료 해시가 기준선과 다르다(라이브러리 판 차이 가능): {item['path']} "
                                f"{_sha256(data)}")
    return warnings


def verify_sources(evalset: dict, root: Path) -> None:
    """평가셋이 적은 파일과 해시가 실제 파일과 같은지 확인한다. 다르면 기준선이 다른 자료를 잰다."""
    for item in source_files(evalset):
        if "derive" in item:
            continue
        if _sha256((root / item["path"]).read_bytes()) != item["sha256"]:
            raise ValueError(f"자료 해시가 평가셋과 다르다: {item['path']}")


async def clean(evalset: dict, dsn: str) -> int:
    owners = list(evalset["users"].values())
    async with await psycopg.AsyncConnection.connect(dsn, autocommit=True) as conn:
        cursor = await conn.execute(
            "DELETE FROM documents WHERE owner_id = ANY(%s) AND %s = ANY(tags)",
            (owners, evalset["tag"]),
        )
        return cursor.rowcount


async def load(evalset: dict, root: Path, dsn: str) -> list[dict]:
    verify_sources(evalset, root)
    tag = evalset["tag"]
    users = evalset["users"]
    outcomes = []
    async with await psycopg.AsyncConnection.connect(dsn, autocommit=True) as conn:
        existing = await (await conn.execute(
            "SELECT count(*) FROM documents WHERE owner_id = ANY(%s) AND %s = ANY(tags)",
            (list(users.values()), tag),
        )).fetchone()
        if existing[0]:
            raise RuntimeError(f"태그 {tag!r} 평가 문서가 이미 {existing[0]}건 있다 — --clean 후 다시 적재한다")
        for source in evalset["sources"]:
            owner = users[source["owner"]]
            path = root / source["path"]
            try:
                document = await create_document(
                    conn, filename=path.name, data=path.read_bytes(), owner_id=owner,
                    title=source["title"], tags=[tag], visibility="private",
                )
            except EmptyExtractedText as error:
                outcomes.append({"title": source["title"], "status": "rejected", "error": str(error)})
                continue
            status = "created"
            stored = source
            if "replaced_by" in source:
                stored = source["replaced_by"]
                replacement = root / stored["path"]
                previous = document["version"]
                document = await replace_original_file(
                    conn, document["id"], user_id=owner, filename=replacement.name,
                    data=replacement.read_bytes(), client_version=previous,
                )
                if document["version"] != previous + 1:
                    raise RuntimeError(f"원본 교체가 다음 텍스트 버전을 만들지 않았다: {source['title']} "
                                       f"v{previous} → v{document['version']}")
                status = "replaced"
            # autocommit 연결의 트랜잭션 밖 SELECT는 OpenProxy가 Replica로 보내 방금 만든 문서를 못 볼 수 있다(ADR-010)
            async with conn.transaction():
                original = await get_original_file(conn, document["id"], user_id=owner)
            if original["sha256"] != _sha256((root / stored["path"]).read_bytes()):
                raise RuntimeError(f"보관된 원본이 넣은 파일과 다르다: {source['title']}")
            outcomes.append({
                "title": source["title"], "status": status, "id": str(document["id"]),
                "version": document["version"], "extraction_status": document["extraction_status"],
                "original_sha256": original["sha256"],
            })
    return outcomes


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n", 1)[0])
    parser.add_argument("evalset", type=Path)
    parser.add_argument("--clean", action="store_true", help="이 평가셋의 태그·계정 문서를 지우고 끝낸다")
    args = parser.parse_args()
    evalset = json.loads(args.evalset.read_text())
    dsn = get_settings().database_url
    if args.clean:
        print(f"지운 평가 문서: {asyncio.run(clean(evalset, dsn))}")
        return
    for warning in ensure_sources(evalset, ROOT):
        print(f"경고: {warning}")
    for outcome in asyncio.run(load(evalset, ROOT, dsn)):
        print(json.dumps(outcome, ensure_ascii=False))


if __name__ == "__main__":
    main()
