"""업로드 파일 바이트에서 추출 텍스트를 만드는 순수 함수.

DB·네트워크를 건드리지 않는다. 일반 추출·래스터화는 메모리에서 처리하고,
OCR은 로컬 tesseract를 호출한다(pytesseract가 임시 파일을 관리한다). 반환된 추출
텍스트는 호출부가 저장한다. 이 모듈은 저장하지 않는다 — 원본 보관은
`services/documents.py`가 `document_files`에 한다.

빈 추출 결과의 거부는 이 모듈 밖에서 한다 — 판정은 `services/documents.py`가
`EmptyExtractedText`로 하고, 400 응답 매핑은 `openarchive/main.py`의 예외 핸들러가 한다.
extract_text는 OCR하지 않는다 — 이미지는 빈 문자열, PDF는 텍스트 레이어만, 오피스 문서는 그림을 뺀
텍스트만 반환한다. 호출부가 needs_ocr로 판정해 워커 추출로 넘긴다(ADR-052). ocr_text는 그때만 호출한다.
"""

import io
import re
import struct
import zipfile
import zlib
from collections.abc import Callable, Iterable, Iterator
from pathlib import PurePath
from xml.etree import ElementTree

import olefile
import pypdfium2
import pytesseract
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape
from pptx.shapes.picture import Picture
from pptx.slide import Slide
from pptx.text.text import TextFrame
from pypdf import PdfReader
from pypdf._codecs.adobe_glyphs import adobe_glyphs

SUPPORTED_CONTENT_TYPES: tuple[str, ...] = (
    "pdf", "docx", "txt", "md", "hwp", "hwpx", "xlsx", "pptx", "png", "jpg", "jpeg"
)

IMAGE_CONTENT_TYPES: tuple[str, ...] = ("png", "jpg", "jpeg")
OCR_LANGUAGE = "kor+eng"
# 기본 psm 3은 깨끗한 300dpi 원본의 문단 블록을 누락했다(#135):
# 1,095자 중 879자, CER 0.35 → psm 4에서 0.07.
OCR_CONFIG = "--psm 4"
# 그림뿐인 오피스 문서(#177): 쪽이 없는 형식은 추출 텍스트가 빌 때 이 경로의 그림을 이름 순서대로
# 인식한다. tesseract가 읽는 래스터 그림만 — EMF·WMF 같은 벡터 그림은 건너뛴다.
_OFFICE_MEDIA_DIRS = {"docx": "word/media/", "xlsx": "xl/media/", "hwpx": "BinData/"}
_RASTER_EXTENSIONS = ("png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff")

# 원본을 내려줄 때의 미디어 타입. 업로더가 보낸 Content-Type이나 저장된 값을 믿지 않고
# 이 고정 매핑만 쓴다 — 조작된 값이 그대로 나가면 브라우저가 다르게 해석한다.
# 형식이 늘면 여기에 한 줄씩 더한다.
MEDIA_TYPES: dict[str, str] = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain; charset=utf-8",
    "md": "text/markdown; charset=utf-8",
    "hwp": "application/x-hwp",
    "hwpx": "application/hwp+zip",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
FALLBACK_MEDIA_TYPE = "application/octet-stream"


def media_type_for(filename: str) -> str:
    """파일명 확장자로 원본 응답의 미디어 타입을 고른다. 모르는 확장자는 octet-stream."""
    return MEDIA_TYPES.get(
        PurePath(filename).suffix.removeprefix(".").lower(), FALLBACK_MEDIA_TYPE
    )


class UnsupportedFileType(ValueError):
    """지원하지 않는 확장자. 사용자에게 보일 메시지를 담는다."""


class TextDecodeError(ValueError):
    """txt/md를 UTF-8로 읽지 못했다."""


class EncryptedDocument(ValueError):
    """암호·배포용 HWP처럼 본문이 암호화돼 읽을 수 없다. 사용자에게 보일 메시지를 담는다."""


def detect_content_type(filename: str) -> str:
    """파일명 확장자로 문서 유형을 판별한다. 지원 형식이 아니면 UnsupportedFileType."""
    content_type = PurePath(filename).suffix.removeprefix(".").lower()
    if content_type not in SUPPORTED_CONTENT_TYPES:
        raise UnsupportedFileType(f"지원하지 않는 파일 형식입니다: {content_type or '확장자 없음'}")
    return content_type


def extract_text(data: bytes, content_type: str) -> str:
    """파일 바이트에서 텍스트를 추출한다.

    PDF 페이지, DOCX·HWP·HWPX 문단, XLSX 시트, PPTX 슬라이드는 빈 줄로 구분한다. 후속
    청킹이 빈 줄을 문단 경계로 사용하므로 파일 안의 구조가 추출 텍스트에도 남는다.

    Raises:
        UnsupportedFileType: `content_type`이 지원하는 유형이 아닐 때.
        TextDecodeError: TXT 또는 Markdown 바이트가 UTF-8이 아닐 때.
        EncryptedDocument: 암호가 걸렸거나 배포용으로 저장된 HWP일 때.
        ValueError: 바이너리 문서 바이트가 비었거나 손상되었을 때.
    """
    if content_type not in SUPPORTED_CONTENT_TYPES:
        raise UnsupportedFileType(f"지원하지 않는 파일 형식입니다: {content_type}")

    if content_type in ("txt", "md"):
        try:
            return data.decode("utf-8")
        except UnicodeDecodeError as error:
            raise TextDecodeError("텍스트 파일은 UTF-8 인코딩이어야 합니다.") from error

    try:
        if content_type in IMAGE_CONTENT_TYPES:
            # OCR은 워커가 한다(ADR-052). 여기서는 끝까지 디코드되는지만 본다 — 못 읽는 파일은
            # 워커가 몇 번을 다시 읽어도 같으므로 다른 형식처럼 업로드에서 거부한다.
            with Image.open(io.BytesIO(data)) as image:
                image.load()
            return ""

        if content_type == "pdf":
            return "\n\n".join(_pdf_page_texts(data))

        if content_type == "hwp":
            return _join_paragraphs(_hwp_paragraphs(data))

        if content_type == "hwpx":
            return _join_paragraphs(_hwpx_paragraphs(data))

        if content_type == "xlsx":
            return _join_paragraphs(_xlsx_sheets(data))

        if content_type == "pptx":
            return _join_paragraphs(_pptx_slides(data))

        document = Document(io.BytesIO(data))
        return "\n\n".join(paragraph.text for paragraph in document.paragraphs)
    except EncryptedDocument:
        raise
    except Exception as error:
        raise ValueError(f"{content_type.upper()} 파일을 읽을 수 없습니다.") from error


def needs_ocr(content_type: str, data: bytes) -> bool:
    """이미지는 항상, PDF는 텍스트 레이어를 글자로 읽을 수 없는 쪽이 하나라도 있을 때 OCR한다.

    읽을 수 없는 쪽은 텍스트 레이어가 빈 쪽과 글자 정보가 깨진 쪽이다(`_pdf_ocr_pages`).
    문서 전체의 텍스트로 판정하면 텍스트 쪽과 스캔 쪽이 섞인 PDF의 스캔 쪽이 빈 채로 남는다.
    오피스 문서도 같은 기준이다(#177) — PPTX는 텍스트 없이 그림만 있는 슬라이드가 있을 때,
    쪽이 없는 DOCX·XLSX·HWPX는 추출 텍스트가 비고 그림이 있을 때 OCR한다.
    호출부는 extract_text가 읽을 수 있음을 확인한 뒤에 부른다.
    """
    if content_type in IMAGE_CONTENT_TYPES:
        return True
    if content_type == "pptx":
        return any(
            not _pptx_slide_has_text(slide) and any(_pptx_pictures(slide.shapes))
            for slide in Presentation(io.BytesIO(data)).slides
        )
    if content_type in _OFFICE_MEDIA_DIRS:
        return bool(_office_pictures(data, content_type)) and not extract_text(
            data, content_type
        ).strip()
    if content_type != "pdf":
        return False
    return any(ocr for _, ocr in _pdf_ocr_pages(PdfReader(io.BytesIO(data)), data))


def _pdf_page_texts(data: bytes) -> list[str]:
    return _pdf_layer_texts(PdfReader(io.BytesIO(data)), data)


def _pdf_layer_texts(reader: PdfReader, data: bytes) -> list[str]:
    """쪽마다 텍스트 레이어 — pypdf 텍스트에 pdfium이 본 같은 줄 간격을 공백으로 더한다(#176).

    pypdf는 줄 끝에서 낱말 중간에 바뀐 줄을 바르게 잇지만, 덩어리마다 그래픽 상태를 따로 여닫는
    PDF(한글 오피스 출력의 표 셀)에서는 셀 사이 간격을 버려 `16,8155,209`처럼 붙인다. pdfium은
    반대로 셀 간격은 지키지만 줄마다 줄바꿈을 넣어 `판⏎단하는`처럼 낱말을 끊는다. 그래서 pypdf를
    기준으로 두고 pdfium의 같은 줄 공백만 빌린다. pdfium이 열지 못하는 파일은 pypdf 텍스트 그대로다.
    """
    texts = [page.extract_text() or "" for page in reader.pages]
    try:
        with pypdfium2.PdfDocument(data) as document:
            for index in range(min(len(document), len(texts))):
                page = document[index]
                try:
                    textpage = page.get_textpage()
                    try:
                        texts[index] = _add_gap_spaces(texts[index], textpage.get_text_range())
                    finally:
                        textpage.close()
                finally:
                    page.close()
    except pypdfium2.PdfiumError:
        pass
    return texts


def _add_gap_spaces(base: str, other: str) -> str:
    """`base`가 붙여 쓴 글자 사이 중 `other`가 같은 줄에서 띄운 자리에만 공백을 넣는다.

    공백을 빼거나 줄바꿈을 더하지 않는다 — `base`의 구분은 그대로 남는다. 공백이 아닌 글자의 순서가
    다르면(읽기 순서가 다른 쪽, 글자 정보가 깨진 쪽) 대응을 믿을 수 없으므로 `base`를 그대로 돌려준다.
    정상 PDF 4,661쪽 중 98%가 같은 순서였다.
    """
    base_chars = [match.start() for match in re.finditer(r"\S", base)]
    other_chars = list(re.finditer(r"\S", other))
    if [base[index] for index in base_chars] != [match[0] for match in other_chars]:
        return base
    parts = []
    start = 0
    for position in range(len(base_chars) - 1):
        end = base_chars[position] + 1
        gap = other[other_chars[position].end() : other_chars[position + 1].start()]
        if end == base_chars[position + 1] and gap and "\n" not in gap and "\r" not in gap:
            parts.append(base[start:end])
            start = end
    parts.append(base[start:])
    return " ".join(parts)


def _pdf_ocr_pages(reader: PdfReader, data: bytes) -> list[tuple[str, bool]]:
    """쪽마다 (레이어 텍스트, OCR할지) — 텍스트 레이어가 비었거나 글자 정보가 깨진 글꼴을 쓰면 OCR한다(#191).

    깨짐은 추출한 텍스트가 아니라 글꼴 구조로 판정한다. 깨진 글자는 글꼴마다 모양이 달라
    (부분 집합 글꼴은 제어 문자, 통 글꼴은 여러 문자 체계, Type3는 ASCII 글리프 이름) 문자 비율로는
    놓친다. 정상 PDF 4,705쪽에서 이 판정에 걸린 쪽은 없었다.
    """
    return [
        (text, not text.strip() or _has_unmapped_font(page.get("/Resources"), set()))
        for page, text in zip(reader.pages, _pdf_layer_texts(reader, data), strict=True)
    ]


def _has_unmapped_font(resources, seen: set[int]) -> bool:
    """ToUnicode가 없고 pypdf가 글자로 옮길 수 없는 글꼴이 이 자원(폼 XObject 포함)에 있는가.

    - Type0·`Identity-H`/`Identity-V`: 글리프 번호가 그대로 글자로 나온다.
    - `/Differences`에 Adobe 글리프 목록에 없는 이름: pypdf가 `/HFT1` 같은 이름을 그대로 낸다.
    그 밖의 ToUnicode 없는 글꼴(WinAnsi·MacRoman 단순 글꼴 등)은 인코딩으로 정상 추출된다.
    """
    if resources is None:
        return False
    resources = resources.get_object()
    if id(resources) in seen:
        return False
    seen.add(id(resources))
    for font in (resources.get("/Font") or {}).values():
        font = font.get_object()
        if "/ToUnicode" in font:
            continue
        encoding = font.get("/Encoding")
        encoding = encoding.get_object() if encoding is not None else None
        if font.get("/Subtype") == "/Type0" and encoding in ("/Identity-H", "/Identity-V"):
            return True
        if isinstance(encoding, dict) and any(
            isinstance(name, str) and name not in adobe_glyphs
            for name in encoding.get("/Differences", [])
        ):
            return True
    return any(
        xobject.get_object().get("/Subtype") == "/Form"
        and _has_unmapped_font(xobject.get_object().get("/Resources"), seen)
        for xobject in (resources.get("/XObject") or {}).values()
    )


def ocr_text(data: bytes, content_type: str) -> str:
    """로컬 OCR로 텍스트를 읽는다. 파일 오류는 ValueError, 설치 오류는 그대로 전파한다.

    PDF는 텍스트 레이어가 빈 쪽과 글자 정보가 깨진 쪽만 인식하고, 나머지 쪽은 extract_text와 같은
    레이어 텍스트를 쓴다. 쪽 순서는 원본 그대로다. PPTX는 텍스트 없는 슬라이드의 그림만 인식하고
    나머지 슬라이드는 extract_text와 같은 텍스트를 쓴다. DOCX·XLSX·HWPX는 문서 안의 그림을 인식한다.
    """
    if content_type not in (*IMAGE_CONTENT_TYPES, "pdf", "pptx", *_OFFICE_MEDIA_DIRS):
        raise UnsupportedFileType(f"지원하지 않는 OCR 파일 형식입니다: {content_type}")

    try:
        if content_type in IMAGE_CONTENT_TYPES:
            return _ocr_image(data)

        if content_type == "pptx":
            return _join_paragraphs(_pptx_slides(data, read_picture=_ocr_image))

        if content_type in _OFFICE_MEDIA_DIRS:
            return _join_paragraphs(
                _ocr_image(picture) for picture in _office_pictures(data, content_type)
            )

        pages = _pdf_ocr_pages(PdfReader(io.BytesIO(data)), data)
        texts = [text for text, _ in pages]
        ocr_pages = [ocr for _, ocr in pages]
        with pypdfium2.PdfDocument(data) as document:
            for index in range(len(document)):
                if not ocr_pages[index]:
                    continue
                page = document[index]
                try:
                    bitmap = page.render(scale=300 / 72)
                    try:
                        with bitmap.to_pil() as image:
                            texts[index] = pytesseract.image_to_string(
                                image, lang=OCR_LANGUAGE, config=OCR_CONFIG
                            )
                    finally:
                        bitmap.close()
                finally:
                    page.close()
        return "\n\n".join(texts)
    except pytesseract.TesseractNotFoundError:
        raise
    except pytesseract.TesseractError as error:
        if any(message in str(error) for message in (
            "Error opening data file", "Failed loading language",
            "Could not initialize tesseract", "couldn't load any languages",
        )):
            raise
        raise ValueError(f"{content_type.upper()} 파일을 읽을 수 없습니다.") from error
    except Exception as error:
        raise ValueError(f"{content_type.upper()} 파일을 읽을 수 없습니다.") from error


def _ocr_image(data: bytes) -> str:
    with Image.open(io.BytesIO(data)) as source, ImageOps.exif_transpose(source) as image:
        return pytesseract.image_to_string(image, lang=OCR_LANGUAGE, config=OCR_CONFIG)


def _office_pictures(data: bytes, content_type: str) -> list[bytes]:
    """DOCX·XLSX·HWPX ZIP 안의 래스터 그림을 이름 순서(`image2`가 `image10` 앞)대로 꺼낸다."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = [
            name for name in archive.namelist()
            if name.startswith(_OFFICE_MEDIA_DIRS[content_type])
            and name.rsplit(".", 1)[-1].lower() in _RASTER_EXTENSIONS
        ]
        names.sort(key=lambda name: [int(part) if part.isdigit() else part
                                     for part in re.split(r"(\d+)", name)])
        return [archive.read(name) for name in names]


def _join_paragraphs(paragraphs: Iterable[str]) -> str:
    return "\n\n".join(paragraph for paragraph in paragraphs if paragraph.strip())


# HWP·HWPX 문단은 본문 순서대로 낸다 — 표 셀·머리말 안의 문단도 그것을 담은 문단 바로 뒤에
# 온다. 두 형식이 같은 문서에서 같은 텍스트를 내도록 특수 문자를 같은 값으로 옮긴다
# (탭·줄바꿈·묶음 빈칸·고정폭 빈칸·하이픈).
#
# 본 제품은 한컴의 HWP 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다.
# (HWP 읽기는 한컴이 공개한 「한글 문서 파일 형식 5.0」을 따른다 — 공개 조건이 소스에 이 문장을
# 적으라고 한다, ADR-059 결정 1)

_HWP_SIGNATURE = b"HWP Document File"
_HWP_COMPRESSED, _HWP_PASSWORD, _HWP_DISTRIBUTION = 0x01, 0x02, 0x04
_HWPTAG_PARA_TEXT = 0x10 + 51
# 문단 텍스트 안 제어 문자 중 인라인·확장 제어는 코드 뒤에 7 WCHAR를 더 차지한다
# (표·그림·각주 자리 등). 나머지 문자 제어는 1 WCHAR다. 한글 문서 파일 형식 5.0 표 6.
_HWP_WIDE_CONTROLS = frozenset({1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23})
_HWP_CONTROL_TEXT = {9: "\t", 10: "\n", 24: "-", 30: " ", 31: " "}


def _hwp_paragraphs(data: bytes) -> Iterator[str]:
    """HWP 5.0 OLE 복합 문서의 `BodyText/Section*` 스트림에서 문단 텍스트 레코드를 읽는다."""
    with olefile.OleFileIO(data) as ole:
        header = ole.openstream("FileHeader").read()
        if not header.startswith(_HWP_SIGNATURE):
            raise ValueError("HWP 서명이 없습니다.")
        flags = int.from_bytes(header[36:40], "little")
        if flags & _HWP_PASSWORD:
            raise EncryptedDocument("암호가 걸린 HWP 문서는 읽을 수 없습니다. 암호를 풀고 다시 저장해 주세요.")
        if flags & _HWP_DISTRIBUTION:
            raise EncryptedDocument("배포용 HWP 문서는 읽을 수 없습니다. 일반 문서로 다시 저장해 주세요.")

        sections = sorted(
            (entry for entry in ole.listdir() if entry[0] == "BodyText" and len(entry) == 2),
            key=lambda entry: int(entry[1].removeprefix("Section")),
        )
        if not sections:
            raise ValueError("HWP 본문 섹션이 없습니다.")
        for entry in sections:
            stream = ole.openstream(entry).read()
            if flags & _HWP_COMPRESSED:
                stream = zlib.decompress(stream, -15)
            yield from _hwp_section_paragraphs(stream)


def _hwp_section_paragraphs(stream: bytes) -> Iterator[str]:
    """섹션 스트림의 레코드(헤더 4바이트: 태그 10비트·레벨 10비트·크기 12비트)를 순회한다."""
    position = 0
    while position < len(stream):
        (header,) = struct.unpack_from("<I", stream, position)
        position += 4
        tag, size = header & 0x3FF, header >> 20
        if size == 0xFFF:  # 4095바이트 이상이면 실제 크기가 뒤따른다.
            (size,) = struct.unpack_from("<I", stream, position)
            position += 4
        if tag == _HWPTAG_PARA_TEXT:
            yield _hwp_paragraph_text(stream[position : position + size])
        position += size


def _hwp_paragraph_text(record: bytes) -> str:
    units = struct.unpack(f"<{len(record) // 2}H", record[: len(record) // 2 * 2])
    text: list[int] = []
    index = 0
    while index < len(units):
        unit = units[index]
        if unit >= 32:
            text.append(unit)
        elif unit in _HWP_CONTROL_TEXT:
            text.append(ord(_HWP_CONTROL_TEXT[unit]))
        if unit in _HWP_WIDE_CONTROLS:
            index += 8
        else:
            index += 1
    # 서로게이트 쌍이 두 단위로 나뉘어 있으므로 문자 단위가 아니라 UTF-16으로 되돌린다.
    return struct.pack(f"<{len(text)}H", *text).decode("utf-16-le", errors="replace")


_HWPX_PARAGRAPH = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"
_HWPX_SECTION = re.compile(r"Contents/section(\d+)\.xml")
_HWPX_CONTROL_TEXT = {"tab": "\t", "lineBreak": "\n", "nbSpace": " ", "fwSpace": " ", "hyphen": "-"}


def _hwpx_paragraphs(data: bytes) -> Iterator[str]:
    """HWPX(OWPML) ZIP의 `Contents/section*.xml`에서 `hp:p` 문단을 문서 순서대로 읽는다."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        sections = sorted(
            (match for name in archive.namelist() if (match := _HWPX_SECTION.fullmatch(name))),
            key=lambda match: int(match[1]),
        )
        if not sections:
            raise ValueError("HWPX 섹션이 없습니다.")
        for section in sections:
            root = ElementTree.fromstring(archive.read(section[0]))
            for paragraph in root.iter(f"{_HWPX_PARAGRAPH}p"):
                # 자기 run의 글자만 모은다 — 표 안 문단은 iter가 뒤이어 따로 낸다.
                yield "".join(
                    _hwpx_text(text)
                    for run in paragraph.findall(f"{_HWPX_PARAGRAPH}run")
                    for text in run.findall(f"{_HWPX_PARAGRAPH}t")
                )


def _hwpx_text(element: ElementTree.Element) -> str:
    parts = [element.text or ""]
    for child in element:
        parts.append(_HWPX_CONTROL_TEXT.get(child.tag.removeprefix(_HWPX_PARAGRAPH), ""))
        parts.append(child.tail or "")
    return "".join(parts)


# XLSX 시트와 PPTX 슬라이드는 한 덩어리로 내고 덩어리 사이를 빈 줄로 둔다. 표는 두 형식 모두
# 행마다 한 줄, 셀은 탭으로 구분한다.


def _xlsx_sheets(data: bytes) -> Iterator[str]:
    """시트마다 시트명을 첫 줄로, 값이 있는 행을 뒤 줄로 낸다. 값 없는 시트는 내지 않는다.

    수식 셀은 파일에 캐시된 계산값을 쓴다(`data_only`). 계산 엔진 없이 쓴 파일처럼 값이
    캐시되지 않았으면 빈 칸이 된다 — 수식을 계산하지도, 수식 문자열을 본문에 넣지도 않는다.
    """
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = [
                _tab_row("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if row.strip()]
            if rows:
                yield "\n".join([sheet.title, *rows])
    finally:
        workbook.close()


def _pptx_slides(data: bytes, read_picture: Callable[[bytes], str] | None = None) -> Iterator[str]:
    """슬라이드마다 도형 텍스트를 순서대로, 그 뒤에 발표자 노트를 낸다.

    `read_picture`를 주면 도형 텍스트가 없는 슬라이드는 그림을 그것으로 읽어 도형 텍스트 자리에 둔다(#177).
    """
    for slide in Presentation(io.BytesIO(data)).slides:
        lines = list(_pptx_shape_texts(slide.shapes))
        if read_picture is not None and not any(line.strip() for line in lines):
            lines = [read_picture(picture) for picture in _pptx_pictures(slide.shapes)]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            lines.append(_pptx_frame_text(slide.notes_slide.notes_text_frame))
        yield "\n".join(line for line in lines if line.strip())


def _pptx_shape_texts(shapes: Iterable[BaseShape]) -> Iterator[str]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pptx_shape_texts(shape.shapes)
        elif shape.has_text_frame:
            yield _pptx_frame_text(shape.text_frame)
        elif shape.has_table:
            for row in shape.table.rows:
                yield _tab_row(_pptx_frame_text(cell.text_frame) for cell in row.cells)


def _pptx_slide_has_text(slide: Slide) -> bool:
    return any(line.strip() for line in _pptx_shape_texts(slide.shapes))


def _pptx_pictures(shapes: Iterable[BaseShape]) -> Iterator[bytes]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pptx_pictures(shape.shapes)
        # 외부 링크만 건 그림은 파일 안에 없어 읽을 수 없다(`image`가 ValueError)
        elif (
            isinstance(shape, Picture)
            and shape._pic.blip_rId is not None
            and shape.image.ext.lower() in _RASTER_EXTENSIONS
        ):
            yield shape.image.blob


def _pptx_frame_text(frame: TextFrame) -> str:
    # python-pptx는 문단 안 줄바꿈(`a:br`)을 수직 탭으로 돌려준다.
    return frame.text.replace("\v", "\n")


def _tab_row(cells: Iterable[str]) -> str:
    # 앞·가운데 빈 칸은 탭으로 남겨 열 위치를 지키고, 줄 끝 빈 칸은 버린다.
    return "\t".join(cells).rstrip("\t")
