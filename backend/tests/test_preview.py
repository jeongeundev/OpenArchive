import io
import subprocess
from pathlib import Path

import psycopg
import pytest
from pypdf import PdfReader, PdfWriter

from openarchive.services import preview

FIXTURES = Path(__file__).parent / "fixtures"


def blank_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    result = io.BytesIO()
    writer.write(result)
    return result.getvalue()


async def test_extensions_match_database(migrated_db):
    extensions = preview.PREVIEW_CONVERTIBLE_EXTENSIONS | {"pdf", "png", "txt", "md", "jpg", "jpeg"}
    async with await psycopg.AsyncConnection.connect(migrated_db) as conn:
        for ext in extensions:
            for name in (f"input.{ext}", f"input.{ext.upper()}"):
                cursor = await conn.execute("SELECT preview_convertible(%s)", (name,))
                assert (await cursor.fetchone())[0] == (
                    ext in preview.PREVIEW_CONVERTIBLE_EXTENSIONS
                )


@pytest.mark.parametrize(
    "text, expected",
    [
        ("", False),
        ("abcㄱㅏ", False),
        ("\uabff", False),
        ("가", True),
        ("힣", True),
        ("\ud7a4", False),
    ],
)
def test_contains_hangul(text, expected):
    assert preview.contains_hangul(text) is expected


def test_pdf_text_layer():
    assert not preview.pdf_has_hangul(blank_pdf())
    assert preview.pdf_has_hangul((FIXTURES / "mixed_tax_pages.pdf").read_bytes())


@pytest.mark.parametrize(
    "extension, setting", [("hwp", "PREVIEW_RHWP_BIN"), ("docx", "PREVIEW_SOFFICE_BIN")]
)
def test_missing_converter(monkeypatch, extension, setting):
    monkeypatch.setenv(setting, "/nonexistent/openarchive-converter")
    with pytest.raises(preview.ConverterUnavailable):
        preview.convert_to_pdf(b"input", f"input.{extension}", document_text="")


def test_missing_sandbox_never_calls_converter(monkeypatch, tmp_path):
    marker = tmp_path / "called"
    converter = tmp_path / "converter"
    converter.write_text(f"#!/bin/sh\ntouch '{marker}'\n")
    converter.chmod(0o755)
    monkeypatch.setenv("PREVIEW_RHWP_BIN", str(converter))
    monkeypatch.setenv("PREVIEW_BWRAP_BIN", "/nonexistent/bwrap")
    with pytest.raises(preview.ConverterUnavailable):
        preview.convert_to_pdf(b"input", "input.hwp", document_text="")
    assert not marker.exists()


def test_unsupported_extension():
    with pytest.raises(ValueError):
        preview.convert_to_pdf(b"input", "input.pdf", document_text="")


@pytest.mark.parametrize("name", ["../한글 ; -o bad.HWP", "a b;bad.hwpx", "../-o.docx"])
def test_fixed_input_arguments(monkeypatch, name):
    monkeypatch.setattr(preview, "_executable", lambda name: name)
    monkeypatch.setattr(preview, "check_converter", lambda ext: "/usr/bin/converter")
    monkeypatch.setattr(preview, "check_fonts", lambda: None)
    monkeypatch.setattr(preview, "check_sandbox", lambda args: None)
    monkeypatch.setattr(preview, "_sandbox_args", lambda *args: ["sandbox"])
    calls = []

    def run(args, timeout):
        calls.append(args)
        if "export-pdf" in args:
            output = Path(args[args.index("-o") + 1])
        else:
            output = Path(args[args.index("--outdir") + 1]) / "input.pdf"
        output.write_bytes(blank_pdf())

    monkeypatch.setattr(preview, "_run", run)
    assert preview.convert_to_pdf(b"input", name, document_text="") == blank_pdf()
    assert all(name not in arg for arg in calls[0])
    assert Path(calls[0][-3] if "export-pdf" in calls[0] else calls[0][-1]).name == (
        "input." + name.rsplit(".", 1)[1].lower()
    )


@pytest.mark.parametrize("output", [None, b"invalid", blank_pdf()])
def test_bad_render_is_permanent(monkeypatch, output):
    monkeypatch.setattr(preview, "_executable", lambda name: name)
    monkeypatch.setattr(preview, "check_converter", lambda ext: "/usr/bin/converter")
    monkeypatch.setattr(preview, "check_fonts", lambda: None)
    monkeypatch.setattr(preview, "check_sandbox", lambda args: None)
    monkeypatch.setattr(preview, "_sandbox_args", lambda *args: [])

    def run(args, timeout):
        if output is not None:
            Path(args[-1]).write_bytes(output)

    monkeypatch.setattr(preview, "_run", run)
    with pytest.raises(preview.PreviewRenderFailed):
        preview.convert_to_pdf(b"input", "input.hwp", document_text="한글")


@pytest.mark.converters
@pytest.mark.parametrize(
    "filename",
    [
        "tax_administration.hwp",
        "tax_administration.hwpx",
        "office_budget.xlsx",
        "office_briefing.pptx",
        "input.docx",
    ],
)
def test_real_conversion(filename):
    if filename == "input.docx":
        from docx import Document

        document = Document()
        document.add_paragraph("한글 미리보기 확인")
        buffer = io.BytesIO()
        document.save(buffer)
        data = buffer.getvalue()
    else:
        data = (FIXTURES / filename).read_bytes()
    pdf = preview.convert_to_pdf(data, filename, document_text="한글")
    assert len(PdfReader(io.BytesIO(pdf)).pages) >= 1
    assert preview.pdf_has_hangul(pdf)


@pytest.mark.converters
def test_timeout_kills_children(monkeypatch, tmp_path):
    import time
    from uuid import uuid4

    child_marker = f"preview-child-{uuid4()}"
    converter = tmp_path / "sleep-converter"
    converter.write_text(f"#!/bin/sh\n/bin/sh -c 'sleep 60' {child_marker} &\nwait\n")
    converter.chmod(0o755)
    monkeypatch.setenv("PREVIEW_RHWP_BIN", str(converter))
    monkeypatch.setenv("PREVIEW_TIMEOUT_SECONDS", "0.5")
    started = time.monotonic()
    with pytest.raises(subprocess.TimeoutExpired):
        preview.convert_to_pdf(b"input", "input.hwp", document_text="")
    assert time.monotonic() - started < 10
    processes = subprocess.check_output(["ps", "-eo", "args"], text=True)
    assert child_marker not in processes
    assert "sleep 60" not in processes


def linked_docx(target):
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml.ns import qn
    from PIL import Image

    image = io.BytesIO()
    Image.new("RGB", (91, 73), (231, 17, 89)).save(image, format="PNG")
    document = Document()
    document.add_paragraph("외부 그림 검사")
    picture = document.add_picture(io.BytesIO(image.getvalue()))
    blip = picture._inline.xpath(".//a:blip")[0]
    old = blip.attrib.pop(qn("r:embed"))
    del document.part.rels[old]
    blip.set(
        qn("r:link"), document.part.relate_to(target, RELATIONSHIP_TYPE.IMAGE, is_external=True)
    )
    result = io.BytesIO()
    document.save(result)
    return result.getvalue(), image.getvalue()


def direct_office(data, directory, *, env=None):
    directory.mkdir()
    source = directory / "input.docx"
    source.write_bytes(data)
    subprocess.run(
        [
            preview.check_converter("docx"),
            "--headless",
            "--norestore",
            f"-env:UserInstallation={(directory / 'profile').as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(directory),
            str(source),
        ],
        check=True,
        capture_output=True,
        timeout=60,
        env=env,
    )
    return (directory / "input.pdf").read_bytes()


def secret_in_pdf(pdf):
    for page in PdfReader(io.BytesIO(pdf)).pages:
        for image in page.images:
            decoded = image.image.convert("RGB")
            if decoded.size == (91, 73) and set(decoded.getdata()) == {(231, 17, 89)}:
                return True
    return False


@pytest.mark.converters
def test_network_isolation_with_positive_control(tmp_path):
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    requests = []
    image_data = linked_docx("http://example.invalid/image.png")[1]

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path)
            self.send_response(200)
            self.send_header("Content-Type", "image/png")
            self.end_headers()
            self.wfile.write(image_data)

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        data, _ = linked_docx(f"http://127.0.0.1:{server.server_port}/secret.png")
        pdf = preview.convert_to_pdf(data, "input.docx", document_text="외부 그림 검사")
        assert len(PdfReader(io.BytesIO(pdf)).pages) >= 1
        assert requests == []
        direct_office(data, tmp_path / "control")
        assert len(requests) >= 1
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


@pytest.mark.converters
def test_local_file_isolation_with_positive_control(tmp_path):
    secret = tmp_path / "secret.png"
    data, image = linked_docx(secret.as_uri())
    secret.write_bytes(image)
    isolated = preview.convert_to_pdf(data, "input.docx", document_text="")
    assert not secret_in_pdf(isolated)
    assert secret_in_pdf(direct_office(data, tmp_path / "control"))


@pytest.mark.converters
def test_host_home_is_hidden(tmp_path):
    import tempfile

    with tempfile.NamedTemporaryFile(dir=Path.home()) as secret:
        source = tmp_path / "input.hwp"
        source.touch()
        output = tmp_path / "output"
        output.mkdir()
        sandbox = preview._sandbox_args(preview.check_converter("hwp"), source, output)
        subprocess.run(
            [*sandbox, "/bin/sh", "-c", 'test ! -e "$1"', "sh", secret.name],
            check=True,
            capture_output=True,
            timeout=10,
        )


@pytest.mark.converters
def test_no_fonts_is_unavailable_and_office_keeps_text(monkeypatch, tmp_path):
    import os

    config = tmp_path / "fonts.conf"
    config.write_text(
        '<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">'
        "<fontconfig><dir>/nonexistent-openarchive-fonts</dir></fontconfig>"
    )
    monkeypatch.setenv("FONTCONFIG_FILE", str(config))
    from docx import Document

    document = Document()
    document.add_paragraph("한글 글꼴 없는 출력 검사")
    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    with pytest.raises(preview.ConverterUnavailable, match="한글 글꼴이 없습니다"):
        preview.convert_to_pdf(data, "input.docx", document_text="한글")
    # 글꼴이 아예 없으면 Ubuntu LO 24는 PDF 생성 전 종료한다.
    with pytest.raises(subprocess.CalledProcessError) as caught:
        direct_office(data, tmp_path / "no-font-control", env=os.environ.copy())
    assert b"No fonts could be found" in caught.value.stderr
    assert not (tmp_path / "no-font-control" / "input.pdf").exists()

    # 한글만 못 그리는 환경도 측정한다. 시스템의 비한글 글꼴 하나만 노출한다.
    import shutil

    installed_env = os.environ.copy()
    installed_env.pop("FONTCONFIG_FILE")
    english = subprocess.check_output(["fc-list", ":lang=en", "file"], text=True, env=installed_env)
    korean = subprocess.check_output(["fc-list", ":lang=ko", "file"], text=True, env=installed_env)
    english_files = {line.split(":", 1)[0] for line in english.splitlines()}
    korean_files = {line.split(":", 1)[0] for line in korean.splitlines()}
    latin = sorted(english_files - korean_files)
    assert latin, "측정 대조군에는 비한글 글꼴이 필요합니다"
    fonts = tmp_path / "latin-fonts"
    fonts.mkdir()
    shutil.copyfile(latin[0], fonts / Path(latin[0]).name)
    config.write_text(f"<fontconfig><dir>{fonts}</dir></fontconfig>")
    with pytest.raises(preview.ConverterUnavailable, match="한글 글꼴이 없습니다"):
        preview.convert_to_pdf(data, "input.docx", document_text="한글")
    pdf = direct_office(data, tmp_path / "latin-font-control", env=os.environ.copy())
    assert preview.pdf_has_hangul(pdf)


def test_font_probe_rejects_empty_output(monkeypatch):
    monkeypatch.setattr(preview, "_executable", lambda name: name)
    monkeypatch.setattr(
        subprocess,
        "run",
        lambda *args, **kwargs: subprocess.CompletedProcess(args, 0, stdout=b"", stderr=b""),
    )
    with pytest.raises(preview.ConverterUnavailable, match="한글 글꼴이 없습니다"):
        preview.check_fonts()


def test_sandbox_probe_never_falls_back(monkeypatch):
    def fail(*args):
        raise subprocess.CalledProcessError(1, "bwrap", stderr="namespace denied")

    monkeypatch.setattr(preview, "_run", fail)
    with pytest.raises(preview.ConverterUnavailable, match="namespace denied"):
        preview.check_sandbox(["bwrap"])


def test_nonzero_exit_contains_bounded_stderr():
    import sys

    with pytest.raises(subprocess.CalledProcessError) as caught:
        preview._run(
            [sys.executable, "-c", "import sys; sys.stderr.write('x'*9000+'TAIL'); exit(3)"], 10
        )
    assert len(caught.value.stderr) <= 4096
    assert str(caught.value).endswith("TAIL")


def test_preview_configuration(monkeypatch):
    from pydantic import ValidationError

    from openarchive.config import Settings

    monkeypatch.setenv("PREVIEW_RHWP_BIN", "/opt/rhwp")
    monkeypatch.setenv("PREVIEW_SOFFICE_BIN", "/usr/bin/soffice")
    monkeypatch.setenv("PREVIEW_BWRAP_BIN", "/usr/bin/bwrap")
    monkeypatch.setenv("PREVIEW_TIMEOUT_SECONDS", "1.5")
    settings = Settings(_env_file=None)
    assert (
        settings.preview_rhwp_bin,
        settings.preview_soffice_bin,
        settings.preview_bwrap_bin,
        settings.preview_timeout_seconds,
    ) == ("/opt/rhwp", "/usr/bin/soffice", "/usr/bin/bwrap", 1.5)
    monkeypatch.setenv("PREVIEW_TIMEOUT_SECONDS", "0")
    with pytest.raises(ValidationError):
        Settings(_env_file=None)


@pytest.mark.converters
def test_fixed_name_reaches_real_converter_process(monkeypatch, tmp_path):
    import json

    converter = tmp_path / "argv-converter"
    # /usr/bin/python3도 샌드박스의 시스템 경로 안에서 실행된다.
    converter.write_text(
        "#!/usr/bin/python3\nimport sys,json\nfrom pathlib import Path\n"
        "output=Path(sys.argv[-1])\n"
        f"output.write_bytes({blank_pdf()!r})\n"
        "output.with_suffix('.json').write_text(json.dumps(sys.argv))\n"
    )
    converter.chmod(0o755)
    monkeypatch.setenv("PREVIEW_RHWP_BIN", str(converter))
    original = preview._run
    seen = []

    def record(args, timeout):
        original(args, timeout)
        if "export-pdf" in args:
            seen.extend(json.loads(Path(args[-1]).with_suffix(".json").read_text()))

    monkeypatch.setattr(preview, "_run", record)
    filename = "../한글 ; -o injected.HWP"
    preview.convert_to_pdf(b"input", filename, document_text="")
    assert seen[1] == "export-pdf"
    assert Path(seen[2]).name == "input.hwp"
    assert seen[3] == "-o"
    assert filename not in " ".join(seen)


def test_sandbox_only_binds_system_input_output_and_single_executable(monkeypatch, tmp_path):
    monkeypatch.setattr(preview, "_executable", lambda name: name)
    converter = tmp_path / "private" / "rhwp"
    source = tmp_path / "input.hwp"
    output = tmp_path / "output"
    args = preview._sandbox_args(str(converter), source, output)
    for flag in (
        "--unshare-net",
        "--unshare-user",
        "--unshare-ipc",
        "--unshare-uts",
        "--unshare-pid",
        "--die-with-parent",
        "--new-session",
        "--clearenv",
    ):
        assert flag in args
    readonly = [(args[i + 1], args[i + 2]) for i, arg in enumerate(args) if arg == "--ro-bind"]
    assert (str(converter), "/converter") in readonly
    assert (str(source), str(source)) in readonly
    assert not any(
        src in {"/", str(converter.parent), str(Path.home()), "/var", "/tmp"} for src, _ in readonly
    )
    writable = [(args[i + 1], args[i + 2]) for i, arg in enumerate(args) if arg == "--bind"]
    assert writable == [(str(output), str(output))]
    assert args[args.index("--tmpfs") + 1] == "/tmp"
    assert "--proc" in args and "--dev" in args
